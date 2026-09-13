"""Regenerate data/earnings_overrides.json from the S&P Global quarterly EPS workbook.

Run manually each earnings season (roughly Feb, May, Aug, Nov) after replacing
reference_resources/sp-500-eps-est.xlsx with the latest S&P Global download:

    python3 scripts/build_earnings_overrides.py

Then commit data/earnings_overrides.json and push — the next daily workflow run
picks up the new earnings in the P/E series.
"""

import json
import os
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd

REPO_ROOT = os.path.join(os.path.dirname(__file__), "..")
SOURCE_PATH = os.path.join(REPO_ROOT, "reference_resources", "sp-500-eps-est.xlsx")
OUTPUT_PATH = os.path.join(REPO_ROOT, "data", "earnings_overrides.json")


def read_as_of_cells():
    """The workbook's own as-of cells from SECTOR EPS!B2/B3/B4: when the
    workbook was compiled (B2), the last quarter of confirmed actuals (B3),
    and the first quarter of estimates (B4, informational only)."""
    wb = openpyxl.load_workbook(SOURCE_PATH, data_only=True)
    ws = wb["SECTOR EPS"]
    data_as_of = ws["B2"].value
    actuals_through = ws["B3"].value
    return (
        data_as_of.strftime("%Y-%m-%d") if hasattr(data_as_of, "strftime") else str(data_as_of),
        actuals_through.date() if hasattr(actuals_through, "date") else actuals_through,
    )


def build():
    data_as_of, actuals_through = read_as_of_cells()

    df = pd.read_excel(SOURCE_PATH, sheet_name="QUARTERLY DATA", header=None, engine="openpyxl")
    df = df.iloc[6:, [0, 1, 2]].copy()
    df.columns = ["quarter_end", "operating_eps", "as_reported_eps"]
    df["quarter_end"] = pd.to_datetime(df["quarter_end"], errors="coerce")
    df["as_reported_eps"] = pd.to_numeric(df["as_reported_eps"], errors="coerce")
    df = df[df["quarter_end"].notna() & df["as_reported_eps"].notna()].copy()
    # Cap at the workbook's own actuals_through cell so a preliminary/estimated
    # quarter in the sheet (e.g. the current quarter, still estimates) can never
    # enter as confirmed, and never pollutes the trailing-4-quarter sum.
    df = df[df["quarter_end"].dt.date <= actuals_through].copy()
    df = df.sort_values("quarter_end").reset_index(drop=True)
    df["ttm_eps"] = df["as_reported_eps"].rolling(4).sum().round(2)
    df = df[df["ttm_eps"].notna()].copy()

    entries = []
    for _, row in df.iterrows():
        qe = row["quarter_end"]
        m = qe.month + 1
        y = qe.year
        if m > 12:
            m, y = 1, y + 1
        effective_from = date(y, m, 1)
        # The sheet's quarter_end is the last *trading* day of the quarter
        # (e.g. 2024-06-28), not necessarily the calendar quarter end; derive
        # the true calendar quarter end from effective_from instead.
        calendar_quarter_end = effective_from - timedelta(days=1)
        entries.append(
            {
                "quarter_end": calendar_quarter_end.strftime("%Y-%m-%d"),
                "effective_from": effective_from.strftime("%Y-%m-%d"),
                "quarterly_eps": round(float(row["as_reported_eps"]), 2),
                "ttm_eps": round(float(row["ttm_eps"]), 2),
            }
        )

    out = {
        "description": (
            "S&P 500 quarterly and TTM as-reported (GAAP) EPS. Source: S&P Global "
            "sp-500-eps-est.xlsx. Regenerate this file each earnings season (4x/year) "
            "by replacing reference_resources/sp-500-eps-est.xlsx and running this script."
        ),
        "source": "S&P Global / S&P Dow Jones Indices (sp-500-eps-est.xlsx)",
        "data_as_of": data_as_of,
        "actuals_through": actuals_through.strftime("%Y-%m-%d"),
        "last_updated": datetime.now().strftime("%Y-%m-%d"),
        "entries": entries,
    }
    with open(OUTPUT_PATH, "w") as f:
        json.dump(out, f, indent=2)
    print(f"Wrote {len(entries)} entries, latest: {entries[-1]['quarter_end']}  TTM={entries[-1]['ttm_eps']}")


if __name__ == "__main__":
    build()
